from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path("/Users/mimi/Documents/Code/Github/GSA-work/explore/sythetic")
DOCX_OUT = ROOT / "output/documents/Visa_Cover_Letter_Additional_Financial_Documents_83847122.docx"

# Preset: standard_business_brief.
# Named overrides: A4 embassy-submission page; restrained black/charcoal letter styling;
# memo_masthead adapted to a traditional formal letter.
INK = RGBColor(24, 35, 48)
MUTED = RGBColor(92, 101, 112)
ACCENT = RGBColor(31, 78, 121)


def set_run_font(run, name="Calibri", size=11, bold=None, italic=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeatable_paragraph_format(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body(doc, text, after=7):
    p = doc.add_paragraph()
    set_repeatable_paragraph_format(p, after=after, line=1.10)
    set_run_font(p.add_run(text), size=10.7)
    return p


def set_cellless_bottom_rule(paragraph, color="D7DCE2", size="6", space="4"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)


def configure_list_paragraph(p):
    fmt = p.paragraph_format
    fmt.left_indent = Cm(0.65)
    fmt.first_line_indent = Cm(-0.33)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(2.5)
    fmt.line_spacing = 1.05


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)
section.header_distance = Cm(1.0)
section.footer_distance = Cm(0.9)

# Explicit base style tokens.
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(10.7)
normal.font.color.rgb = INK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.10

for name, size, before, after, color in [
    ("Title", 16, 0, 3, INK),
    ("Heading 1", 14, 14, 6, ACCENT),
    ("Heading 2", 12, 10, 5, ACCENT),
    ("Heading 3", 11, 8, 4, INK),
]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

# Quiet footer for document identification.
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.paragraph_format.space_before = Pt(0)
footer_p.paragraph_format.space_after = Pt(0)
set_run_font(
    footer_p.add_run("Zhiyi Zhu  |  Visa Application 83847122"),
    size=8.5,
    color=MUTED,
)

# Sender masthead.
p = doc.add_paragraph()
set_repeatable_paragraph_format(p, after=0, line=1.0)
set_run_font(p.add_run("ZHIYI ZHU"), size=16, bold=True, color=INK)

p = doc.add_paragraph()
set_repeatable_paragraph_format(p, after=4, line=1.0)
set_run_font(
    p.add_run("zhiyizhu7@gmail.com  |  +49 1744245952"),
    size=9.5,
    color=MUTED,
)
set_cellless_bottom_rule(p)

p = doc.add_paragraph()
set_repeatable_paragraph_format(p, before=7, after=10, line=1.0)
set_run_font(p.add_run("23 July 2026"), size=10.5)

p = doc.add_paragraph()
set_repeatable_paragraph_format(p, after=1, line=1.0)
set_run_font(p.add_run("Visa Officer"), size=10.7, bold=True)

p = doc.add_paragraph()
set_repeatable_paragraph_format(p, after=13, line=1.0)
set_run_font(p.add_run("Visa Section"), size=10.7, color=MUTED)

# Subject block.
p = doc.add_paragraph()
set_repeatable_paragraph_format(p, after=1, line=1.0)
set_run_font(p.add_run("ADDITIONAL FINANCIAL DOCUMENTS"), size=14, bold=True, color=ACCENT)

p = doc.add_paragraph()
set_repeatable_paragraph_format(p, after=12, line=1.0)
set_run_font(p.add_run("Visa Application Number: "), size=10.5, bold=True)
set_run_font(p.add_run("83847122"), size=10.5)

p = doc.add_paragraph()
set_repeatable_paragraph_format(p, after=8, line=1.0)
set_run_font(p.add_run("Dear Visa Officer,"), size=10.7)

add_body(
    doc,
    "I am writing in response to the request for six months of recent "
    "euro-denominated bank statements.",
)

add_body(
    doc,
    "I arrived in Germany on 31 March 2026 and opened my Revolut euro account on "
    "7 April 2026. Because the account has been open for less than six months, six "
    "full months of statements do not yet exist. I have therefore enclosed all "
    "available Revolut statements, covering the period from the account opening "
    "date to the present.",
)

add_body(
    doc,
    "For account security and the prudent management of my savings, I transferred "
    "part of my euro savings to my euro-denominated account with Bank of China. I "
    "have also enclosed the relevant euro account statements for this account.",
)

add_body(
    doc,
    "Both accounts are held solely in my name. Transfers between my Revolut account "
    "and my Bank of China account are therefore transfers between my own personal "
    "accounts and do not constitute financial support or funding from any third party.",
)

add_body(
    doc,
    "I have also enclosed my June 2026 payslip as evidence of my current employment "
    "income in Germany.",
)

add_body(
    doc,
    "I hope the enclosed documents provide sufficient evidence of my financial "
    "circumstances and available funds. Should you require any further information "
    "or documentation, please do not hesitate to contact me.",
)

add_body(doc, "Thank you for your consideration.", after=11)

p = doc.add_paragraph()
set_repeatable_paragraph_format(p, after=10, line=1.0)
set_run_font(p.add_run("Yours faithfully,"), size=10.7)

p = doc.add_paragraph()
set_repeatable_paragraph_format(p, after=11, line=1.0)
set_run_font(p.add_run("Zhiyi Zhu"), size=10.7, bold=True)

p = doc.add_paragraph()
set_repeatable_paragraph_format(p, before=0, after=3, line=1.0)
set_run_font(p.add_run("Enclosures"), size=10.3, bold=True, color=ACCENT)

enclosures = [
    "All available Revolut euro account statements from 7 April 2026 to the present",
    "Relevant Bank of China euro account statements",
    "June 2026 payslip",
]
for text in enclosures:
    p = doc.add_paragraph(style="List Bullet")
    configure_list_paragraph(p)
    set_run_font(p.add_run(text), size=9.5, color=MUTED)

# Core properties.
doc.core_properties.title = "Cover Letter - Additional Financial Documents"
doc.core_properties.subject = "Visa Application 83847122"
doc.core_properties.author = "Zhiyi Zhu"
doc.core_properties.keywords = "visa, financial documents, bank statements"

DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(DOCX_OUT)
print(DOCX_OUT)
